"""Integrantes y equipos (Excel) y resolucion de nombres de usuario. Misma logica que PUBLICACIONES_VENTAS/pub_sync."""
import re
import unicodedata


def persona_norm(v):
    """NormalizarPersona: mayusculas, sin acentos (incluye N), puntuacion como espacio."""
    s = re.sub(r"[\x00-\x1f\x7f]", "", str(v or "")).strip().upper()
    s = "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")
    s = re.sub(r"[.,;:\-_/\\]", " ", s)
    return " ".join(s.split())


def persona_compacta(v):
    """CompactarPersona: 'Juan Pablo Reyes' y 'JUANPABLOREYES' dan la misma llave."""
    return persona_norm(v).replace(" ", "")


def disciplina(modelo):
    """Misma regla que HISTORIAL_PUBLICACIONES_MODELOS (por palabras en el nombre del archivo)."""
    a = str(modelo or "").upper()
    if "ARQ" in a or "ARC" in a:
        return "ARQUITECTURA"
    if "EST" in a or "STR" in a:
        return "ESTRUCTURA"
    if "ELE" in a:
        return "ELECTRICA"
    if "ESP" in a or "SPE" in a:
        return "ESPECIALES"
    if "MEC" in a:
        return "MECANICA"
    if "PLO" in a or "PLU" in a:
        return "PLOMERIA"
    return "SIN DISCIPLINA"


def leer_equipos(ruta, hoja="Integrantes"):
    """-> lista de {integrante, equipo, compacta} (distinta por llave compacta)."""
    from openpyxl import load_workbook
    ws = load_workbook(ruta, data_only=True, read_only=True)[hoja]
    filas = list(ws.iter_rows(values_only=True))
    enc = [str(c or "").strip() for c in filas[0]]
    ie, ii = enc.index("Equipo"), enc.index("Integrante")
    out, vistos = [], set()
    for r in filas[1:]:
        eq = str(r[ie] or "").replace("\xa0", " ").strip()
        it = str(r[ii] or "").replace("\xa0", " ").strip()
        if not eq or not it:
            continue
        k = persona_compacta(it)
        if k in vistos:
            continue
        vistos.add(k)
        out.append({"integrante": it, "equipo": eq, "compacta": k})
    return out


def _variantes_login(integrante):
    """Formas en que un nombre del Excel puede aparecer como usuario de Autodesk:
    cualquier combinacion en orden de 2 o mas de sus palabras, pegadas.
    'Mario Alberto Sanchez Munoz' -> MARIOSANCHEZ, MARIOALBERTOSANCHEZ, MARIOSANCHEZMUNOZ, ..."""
    from itertools import combinations
    palabras = persona_norm(integrante).split()
    out = set()
    for n in range(2, len(palabras) + 1):
        for combo in combinations(palabras, n):
            out.add("".join(combo))
    return out


class Personas:
    """Resuelve un nombre (usuario Revit o de ACC) contra el listado oficial del Excel.

    Prioridad:
      1. Alias de config.json (para casos que no se resuelven solos).
      2. Nombre completo igual al del Excel (sin acentos, mayusculas ni espacios).
      3. Usuario de Autodesk tipo 'mariosanchezgcp': se quita el sufijo 'gcp' y se busca
         el integrante cuyo nombre forme ese usuario (MARIO + SANCHEZ). Solo se acepta si
         coincide con UNA sola persona del listado; si hay duplicados no se adivina.
    Sin coincidencia: nombre original y SIN EQUIPO (queda registrado en no_resueltos)."""

    SUFIJOS_LOGIN = ("GCP",)

    def __init__(self, equipos, alias=None):
        self.por_k = {e["compacta"]: e for e in equipos}
        self.alias = {persona_compacta(k): v for k, v in (alias or {}).items()}
        variantes = {}
        for e in equipos:
            for v in _variantes_login(e["integrante"]):
                variantes.setdefault(v, []).append(e)
        self.por_login = {v: es[0] for v, es in variantes.items() if len(es) == 1}
        self.no_resueltos = set()

    def _por_login(self, k):
        for suf in self.SUFIJOS_LOGIN:
            if k.endswith(suf) and len(k) > len(suf):
                e = self.por_login.get(k[: -len(suf)])
                if e:
                    return e
        return self.por_login.get(k)

    def resolver(self, nombre):
        k = persona_compacta(nombre)
        destino = self.alias.get(k)
        if destino:
            k = persona_compacta(destino)
        e = self.por_k.get(k) or (self._por_login(k) if k else None)
        if e:
            return e["integrante"], e["equipo"]
        if nombre:
            self.no_resueltos.add(nombre)
        return (nombre or None), "SIN EQUIPO"
